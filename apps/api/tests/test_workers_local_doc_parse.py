"""Tests for workers/local/doc_parse.py — local document extraction."""

import base64
import io

import pytest

from workers.local.doc_parse import (
    DocumentParseError,
    parse_bytes,
    run,
)


def _make_docx(body: str) -> bytes:
    """Build a minimal DOCX in-memory using python-docx."""
    from docx import Document

    doc = Document()
    doc.add_paragraph(body)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx(rows: list[list[object]]) -> bytes:
    """Build a minimal XLSX with openpyxl."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestFormatDetection:
    def test_plain_text_recognised(self):
        parsed = parse_bytes(b"hello world", mime_type="text/plain")
        assert parsed.format == "text"
        assert "hello world" in parsed.text

    def test_docx_recognised_from_magic_bytes(self):
        data = _make_docx("The quick brown fox jumps over the lazy dog")
        parsed = parse_bytes(data)
        assert parsed.format == "docx"
        assert "quick brown fox" in parsed.text

    def test_xlsx_recognised_from_magic_bytes(self):
        data = _make_xlsx([["Name", "Age"], ["Alice", 30], ["Bob", 25]])
        parsed = parse_bytes(data)
        assert parsed.format == "xlsx"
        assert "Alice" in parsed.text
        assert "30" in parsed.text

    def test_oversized_document_rejected(self):
        big = b"x" * (21 * 1024 * 1024)  # 21 MB, just over the 20 MB cap
        with pytest.raises(DocumentParseError):
            parse_bytes(big)


class TestRunEntry:
    @pytest.mark.asyncio
    async def test_base64_text(self):
        encoded = base64.b64encode(b"hello world").decode()
        result = await run({"content_base64": encoded})
        assert result["format"] == "text"
        assert "hello world" in result["text"]

    @pytest.mark.asyncio
    async def test_base64_docx(self):
        data = _make_docx("Document body here")
        encoded = base64.b64encode(data).decode()
        result = await run({"content_base64": encoded})
        assert result["format"] == "docx"
        assert "Document body here" in result["text"]

    @pytest.mark.asyncio
    async def test_neither_url_nor_base64_raises(self):
        with pytest.raises(DocumentParseError):
            await run({})

    @pytest.mark.asyncio
    async def test_invalid_base64_raises(self):
        with pytest.raises(DocumentParseError):
            await run({"content_base64": "not!valid!base64!!!"})

    @pytest.mark.asyncio
    async def test_legacy_base64_content_key_also_works(self):
        """Older clients may send ``base64_content`` instead of ``content_base64``."""
        encoded = base64.b64encode(b"legacy key test").decode()
        result = await run({"base64_content": encoded})
        assert "legacy key test" in result["text"]
